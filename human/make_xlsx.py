"""Excel versions of the blank re-check sheets: IDs stored as text, dropdowns for coded fields.
Usage: python3 human/make_xlsx.py  ->  human/1_screening.xlsx, human/2_rows.xlsx"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

H = os.path.dirname(os.path.abspath(__file__))
HEAD = PatternFill("solid", fgColor="DCE6F2")
TODO = PatternFill("solid", fgColor="FFF7D6")

def build(name, lists, widths, todo_from):
    rows = list(csv.reader(open(os.path.join(H, name + ".csv"))))
    wb = Workbook(); ws = wb.active; ws.title = name[2:]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.number_format = "@"                        # text: keeps 2609.20980 intact
            if r == 1:
                cell.font, cell.fill = Font(bold=True), HEAD
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif c > todo_from:
                cell.fill = TODO
        if r > 1 and name.startswith("1_"):
            ws.cell(r, 3).hyperlink = row[2]; ws.cell(r, 3).font = Font(color="1F5FBF", underline="single")
    n = len(rows)
    for col, opts in lists.items():
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n}")
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "B2"
    wb.save(os.path.join(H, name + ".xlsx"))

build("1_screening", {"D": "Y,N", "G": "Y,N", "H": "Y,N", "I": "Y,N", "J": "Y,N", "K": "yes,no,unclear", "L": "Y,N"},
      {"A": 12, "B": 48, "C": 30, "D": 9, "E": 28, "F": 14, "G": 10, "H": 10, "I": 10, "J": 10, "K": 12, "L": 12, "M": 9, "N": 40}, 3)
build("2_rows", {"K": "success,partial", "L": "Y,N", "N": "Y,N"},
      {"A": 12, "B": 5, "C": 32, "D": 5, "E": 30, "F": 14, "G": 6, "H": 6, "I": 6, "J": 6, "K": 10, "L": 10, "M": 22, "N": 10, "O": 20, "P": 40}, 5)
print("xlsx sheets written")
