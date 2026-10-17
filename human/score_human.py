"""Score the blinded human re-check against the final (adjudicated) data.

Usage: python3 human/score_human.py [sheet_dir]     (default: human/)
Writes <sheet_dir>/agreement.json, discrepancies.csv, macros_human.tex. The SHA-256 of both filled sheets is
recorded on the first run; later runs refuse if the sheets changed, so the reported agreement is pre-resolution.
"""
import csv
import glob
import hashlib
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "human")
final = {r["id"]: r for r in (json.load(open(f)) for f in glob.glob(os.path.join(ROOT, "data/extract/*.json")))}

def pick(stem):
    """Use the filled .xlsx if present, else the .csv."""
    x = os.path.join(D, stem + ".xlsx")
    return x if os.path.exists(x) else os.path.join(D, stem + ".csv")


def load(path):
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).active
        vals = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        out = [dict(zip(vals[0], r)) for r in vals[1:]]
    else:
        out = list(csv.DictReader(open(path)))
    out = [r for r in out if (r.get("id") or "").strip()]
    bad = [r["id"] for r in out if r["id"].strip() not in final]
    if bad:
        raise SystemExit(f"Unknown paper IDs in {os.path.basename(path)} (did a spreadsheet turn them into numbers?): {bad}")
    return out


sheets = [pick("1_screening"), pick("2_rows")]
digest = {os.path.basename(p): hashlib.sha256(open(p, "rb").read()).hexdigest() for p in sheets}
lock = os.path.join(D, "sheets.sha256.json")
if os.path.exists(lock) and json.load(open(lock)) != digest:
    raise SystemExit("Sheets changed since first scoring; agreement must be computed on the untouched sheets.")


def yn(v):
    v = (v or "").strip().lower()
    return True if v in ("y", "yes", "true", "1") else False if v in ("n", "no", "false", "0") else None


def num(v):
    v = (v or "").strip()
    return int(float(v)) if v else None


disc = []
scr = load(sheets[0])
blank = [r["id"] for r in scr if yn(r["include (Y/N)"]) is None]
if blank:
    raise SystemExit(f"Screening incomplete for: {blank}")

# ---- screening agreement and Cohen's kappa
pairs = [(final[r["id"]]["include"], yn(r["include (Y/N)"])) for r in scr]
n = len(pairs)
agree = sum(a == b for a, b in pairs)
p1, p2 = sum(a for a, _ in pairs) / n, sum(b for _, b in pairs) / n
pe = p1 * p2 + (1 - p1) * (1 - p2)
kappa = (agree / n - pe) / (1 - pe) if pe < 1 else float("nan")
for r in scr:
    f = final[r["id"]]
    if f["include"] != yn(r["include (Y/N)"]):
        disc.append(dict(id=r["id"], row="", field="include", final=f"{f['include']} ({f['exclude_reason']})",
                         human=f"{r['include (Y/N)']} ({r['reason (I1/I2/I3 failed, or meets I1-I3)']})"))

# ---- reporting fields on papers both include
REP = [("ci_or_errorbars", "ci_or_errorbars (Y/N)"), ("sig_test", "sig_test (Y/N)"),
       ("seeds_reported", "seeds_reported (Y/N)"), ("success_defined", "success_defined (Y/N)"),
       ("randomized_init", "randomized_init (yes/no/unclear)")]
rep = {k: [0, 0] for k, _ in REP}
for r in scr:
    f = final[r["id"]]
    if not (f["include"] and yn(r["include (Y/N)"])):
        continue
    for k, col in REP:
        a = f["reporting"][k]
        b = (r[col] or "").strip().lower() if k == "randomized_init" else yn(r[col])
        if k != "randomized_init":
            a = bool(a)
        rep[k][1] += 1
        rep[k][0] += a == b
        if a != b:
            disc.append(dict(id=r["id"], row="", field="reporting." + k, final=a, human=r[col]))

# ---- numbers
rows = load(sheets[1])
numeric = dict(rows=0, coded=0, binary_rows=0, binary_exact=0, outcome_agree=0, comparator_flagged=0, added=0)
for r in rows:
    if r["row"].strip().lower() == "new":
        numeric["added"] += 1
        disc.append(dict(id=r["id"], row="new", field="missing_row", final="", human=f"{r['task']} | {r['comparator']}"))
        continue
    f = final[r["id"]]
    if not f["include"] or "excluded in step 1" in (r["notes"] or "").lower():
        continue
    c = f["comparisons"][int(r["row"])]
    numeric["rows"] += 1
    h = tuple(num(r[k]) for k in ("nP", "sP", "nX", "sX"))
    if all(x is None for x in h):
        disc.append(dict(id=r["id"], row=r["row"], field="not_coded", final="", human=""))
        continue
    numeric["coded"] += 1
    h_bin = not (r["outcome (success/partial)"] or "").strip().lower().startswith("partial")
    f_bin = c["outcome"] == "success"
    numeric["outcome_agree"] += h_bin == f_bin
    if h_bin != f_bin:
        disc.append(dict(id=r["id"], row=r["row"], field="outcome", final=c["outcome"], human=r["outcome (success/partial)"]))
    fv = (c["nP"], c["sP"], c["nX"], c["sX"])
    if f_bin:
        numeric["binary_rows"] += 1
        numeric["binary_exact"] += fv == h
    if (fv if f_bin else (fv[0], fv[2])) != (h if f_bin else (h[0], h[2])):
        disc.append(dict(id=r["id"], row=r["row"], field="numbers", final=fv, human=h))
    if yn(r["comparator OK? (Y/N)"]) is False:
        numeric["comparator_flagged"] += 1
        disc.append(dict(id=r["id"], row=r["row"], field=f"kind {c['kind']} choice", final=c["baseline_name"],
                         human=r["should be (if not OK)"]))

not_coded = [(d["id"], d["row"]) for d in disc if d["field"] == "not_coded"]
if not_coded:
    raise SystemExit(f"Rows not yet coded (sheets not locked): {not_coded[:10]}{' ...' if len(not_coded) > 10 else ''}")
if not os.path.exists(lock):
    json.dump(digest, open(lock, "w"), indent=1)      # lock only once both sheets are complete

mins = [float(r["minutes spent"]) for r in scr if (r["minutes spent"] or "").strip()]
out = dict(papers=n, include_agree=agree, include_kappa=kappa,
           reporting_agreement={k: (a / t if t else None, t) for k, (a, t) in rep.items()},
           numeric=numeric, discrepancies=len(disc), minutes_total=sum(mins), sheets_sha256=digest)
json.dump(out, open(os.path.join(D, "agreement.json"), "w"), indent=1)
with open(os.path.join(D, "discrepancies.csv"), "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=["id", "row", "field", "final", "human", "resolution"])
    w.writeheader()
    w.writerows(disc)


def pct(x):
    return f"{100 * x:.0f}\\%" if x == x else "--"


M = dict(HumPapers=n, HumIncAgree=agree, HumKappa=f"{kappa:.2f}", HumRows=numeric["coded"],
         HumBinaryRows=numeric["binary_rows"],
         HumExactPct=pct(numeric["binary_exact"] / numeric["binary_rows"] if numeric["binary_rows"] else float("nan")),
         HumCompFlag=numeric["comparator_flagged"], HumAdded=numeric["added"], HumHours=f"{sum(mins) / 60:.0f}")
with open(os.path.join(D, "macros_human.tex"), "w") as fh:
    fh.write("% generated by human/score_human.py -- do not edit by hand\n")
    for k, v in M.items():
        fh.write(f"\\newcommand{{\\{k}}}{{{v}}}\n")
print(json.dumps(out, indent=1))
